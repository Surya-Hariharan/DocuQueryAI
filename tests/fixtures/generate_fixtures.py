"""
Regenerates the small, synthetic fixture files in this directory. Not part
of the test suite itself, and not run automatically — the generated files
are committed to the repo so tests are deterministic without needing this
script or its extra dependency at test time.

No real-world documents are used or required; everything here is
synthetic, minimal content created purely to exercise the ingestion
pipeline.

Requires reportlab (only for this script, not a project dependency):
    pip install reportlab

Run from the repo root:
    python tests/fixtures/generate_fixtures.py
"""
import io
import os

import PyPDF2
import docx
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

FIXTURES = os.path.dirname(os.path.abspath(__file__))
PDF_DIR = os.path.join(FIXTURES, "pdf")
DOCX_DIR = os.path.join(FIXTURES, "docx")
TEXT_DIR = os.path.join(FIXTURES, "text")


def main():
    os.makedirs(PDF_DIR, exist_ok=True)
    os.makedirs(DOCX_DIR, exist_ok=True)
    os.makedirs(TEXT_DIR, exist_ok=True)

    # pdf/normal.pdf — a real, small PDF with genuinely extractable text
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(72, 720, "DocuQueryAI Test Fixture")
    c.drawString(72, 700, "This is a normal, unencrypted PDF with extractable text.")
    c.drawString(72, 680, "It exists purely to exercise the ingestion pipeline in tests.")
    c.showPage()
    c.save()
    with open(os.path.join(PDF_DIR, "normal.pdf"), "wb") as f:
        f.write(buf.getvalue())

    # pdf/empty.pdf — a single blank page, no extractable text
    writer = PyPDF2.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with open(os.path.join(PDF_DIR, "empty.pdf"), "wb") as f:
        f.write(buf.getvalue())

    # pdf/encrypted.pdf — genuinely password-protected (not just permissions-only)
    writer = PyPDF2.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="fixture-password")
    buf = io.BytesIO()
    writer.write(buf)
    with open(os.path.join(PDF_DIR, "encrypted.pdf"), "wb") as f:
        f.write(buf.getvalue())

    # docx/structured.docx — heading + paragraphs + table
    document = docx.Document()
    document.add_heading("Q4 Financial Summary", level=1)
    document.add_paragraph(
        "Revenue for the fourth quarter reached significant growth across all regions, "
        "driven primarily by the widget division's expansion into new markets."
    )
    document.add_paragraph(
        "Operating costs remained flat year over year despite the increased production volume."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "Approved"
    document.save(os.path.join(DOCX_DIR, "structured.docx"))

    # text/short_content.txt — the realistic short-content examples from the audit
    short_examples = [
        "Q4 Financial Summary",
        "Approved",
        "Pending",
        "N/A",
        "Q4",
        "North Region",
        "Project Alpha",
        "Revenue: ₹12.4M",
    ]
    with open(os.path.join(TEXT_DIR, "short_content.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(short_examples) + "\n")

    print("Fixtures written to:", FIXTURES)


if __name__ == "__main__":
    main()
