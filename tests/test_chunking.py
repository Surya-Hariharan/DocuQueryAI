"""
Regression tests for the short-content-drop bug: split_text_into_chunks()
used to silently discard any section whose entire content fell under
MIN_CHUNK_TOKENS, which destroyed real information (headings, table
cells, status labels) rather than merely denoising fragment leftovers.

These import chunking.py directly, which loads the e5-small-v2 tokenizer
at module import time (a real network fetch on first run, or a local HF
cache) — skips cleanly without `transformers` reachable.
"""

import pytest

pytest.importorskip("transformers", reason="chunking.py loads a HF tokenizer at import time")

from src.ingestion.chunking import split_text_into_chunks, count_tokens, MIN_CHUNK_TOKENS


REALISTIC_SHORT_CONTENT = [
    "Q4 Financial Summary",
    "Approved",
    "Pending",
    "N/A",
    "Q4",
    "North Region",
    "Project Alpha",
    "Revenue: ₹12.4M",  # "Revenue: ₹12.4M"
]


class TestShortContentIsPreserved:
    @pytest.mark.parametrize("text", REALISTIC_SHORT_CONTENT)
    def test_short_meaningful_content_survives_intact(self, text):
        assert count_tokens(text) < MIN_CHUNK_TOKENS, (
            f"Fixture assumption broken: {text!r} is no longer under the "
            f"minimum-token threshold ({MIN_CHUNK_TOKENS}) — this test needs "
            "genuinely short input to be meaningful."
        )
        result = split_text_into_chunks(text)
        assert result == [text], f"Expected {text!r} to survive as its own chunk, got {result!r}"

    def test_multiple_distinct_short_values_each_survive_independently(self):
        # Each call is independent (one per section, as the real ingestion
        # pipeline does it) — confirms there's no hidden cross-call state.
        results = [split_text_into_chunks(t) for t in ["Approved", "Pending", "Approved"]]
        assert results == [["Approved"], ["Pending"], ["Approved"]]

    def test_result_is_deterministic(self):
        text = "North Region"
        assert split_text_into_chunks(text) == split_text_into_chunks(text)


class TestEmptyAndWhitespaceStillProduceNoChunks:
    """The fix must not turn 'preserve short content' into 'never filter
    anything' — genuinely empty input has no content to preserve."""

    def test_empty_string(self):
        assert split_text_into_chunks("") == []

    def test_whitespace_only(self):
        assert split_text_into_chunks("   \n\n  \t  ") == []


class TestLongDocumentFragmentFilteringStillWorks:
    """The original purpose of the filter — dropping a noisy trailing
    fragment from a long, multi-chunk document — must still hold. Only the
    lone-short-chunk case changed."""

    def test_a_long_document_producing_multiple_chunks_still_filters_tiny_trailing_fragments(self):
        # A long paragraph engineered to overflow max_tokens multiple times,
        # ending in a deliberately tiny final sentence.
        long_sentence = ("This is a substantive sentence with real content about quarterly "
                          "performance metrics and operational efficiency improvements. ") * 30
        text = long_sentence + "\n\nOk."
        result = split_text_into_chunks(text, max_tokens=100, overlap_tokens=10)

        assert len(result) > 1, "Expected the long input to actually split into multiple chunks"
        # The tiny trailing fragment ("Ok.") should not survive as its own
        # separate chunk when there are other substantial chunks around it.
        assert not any(c.strip() == "Ok." for c in result)


class TestHeadingAndTableCellsSurviveThroughTheFullPipeline:
    """Integration-level check: the bug was found via a real DOCX through
    parse_document(), not just the leaf chunker in isolation."""

    def test_docx_heading_survives_as_its_own_chunk(self):
        pytest.importorskip("docx", reason="python-docx not installed")
        import io
        import docx
        from src.ingestion.pipeline import parse_document

        document = docx.Document()
        document.add_heading("Q4 Financial Summary", level=1)
        document.add_paragraph(
            "Revenue for the fourth quarter reached significant growth across all regions."
        )
        buf = io.BytesIO()
        document.save(buf)

        chunks = parse_document(buf.getvalue(), filename="summary.docx")
        heading_chunks = [c for c in chunks if c.chunk_type == "heading"]
        assert len(heading_chunks) == 1, f"Expected the heading to survive as its own chunk, got chunks={chunks!r}"
        assert heading_chunks[0].text == "Q4 Financial Summary"

    def test_xlsx_short_status_cell_survives(self):
        pytest.importorskip("openpyxl", reason="openpyxl not installed")
        import io
        import openpyxl
        from src.ingestion.pipeline import parse_document

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Approvals"
        ws.append(["Item", "Status"])
        ws.append(["Widget A", "Approved"])
        buf = io.BytesIO()
        wb.save(buf)

        chunks = parse_document(buf.getvalue(), filename="approvals.xlsx")
        assert len(chunks) == 1
        assert "Approved" in chunks[0].text
        assert chunks[0].sheet_name == "Approvals"
