"""
Format-specific parsers. Each implements DocumentParser.parse(bytes) ->
CanonicalDocument, so adding a new format means adding one class and one
registry entry — never touching the ingestion pipeline itself.
"""

import csv
import io
import logging
from abc import ABC, abstractmethod
from typing import Optional

import openpyxl
from docx import Document as DocxDocument

from src.ingestion.document_model import CanonicalDocument, DocumentSection, SectionType
from src.ingestion.chunking import extract_pages_from_pdf

logger = logging.getLogger("parsers")


class DocumentParser(ABC):
    @abstractmethod
    def parse(self, data: bytes) -> CanonicalDocument:
        raise NotImplementedError


class PdfParser(DocumentParser):
    """Wraps the existing PyPDF2-based page extraction — reused, not rewritten."""

    def parse(self, data: bytes) -> CanonicalDocument:
        pages = extract_pages_from_pdf(io.BytesIO(data))
        sections = [
            DocumentSection(
                text=page_text,
                section_type=SectionType.PARAGRAPH,
                page_number=i + 1,
                position=i
            )
            for i, page_text in enumerate(pages)
            if page_text.strip()
        ]
        return CanonicalDocument(sections=sections, source_type="pdf")


class DocxParser(DocumentParser):
    """
    Preserves headings (by style) as their own sections and keeps each
    table's rows grouped as one table section, rather than flattening
    everything into anonymous paragraphs.

    Known limitation: paragraphs and tables are emitted in two passes
    (all paragraphs, then all tables), not fully interleaved in the
    document's original reading order. Good enough for retrieval today;
    revisit only if reading order turns out to matter for citations.
    """

    def parse(self, data: bytes) -> CanonicalDocument:
        doc = DocxDocument(io.BytesIO(data))
        sections = []
        position = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = (para.style.name or "").lower() if para.style else ""
            section_type = SectionType.HEADING if "heading" in style_name else SectionType.PARAGRAPH
            sections.append(DocumentSection(text=text, section_type=section_type, position=position))
            position += 1

        for table in doc.tables:
            row_lines = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    row_lines.append(" | ".join(cells))
            if row_lines:
                sections.append(DocumentSection(
                    text="\n".join(row_lines),
                    section_type=SectionType.TABLE,
                    position=position
                ))
                position += 1

        return CanonicalDocument(sections=sections, source_type="docx")


class XlsxParser(DocumentParser):
    """
    Never flattens a sheet into raw cell-to-text concatenation: every row
    becomes its own section, with each value paired with its column header
    ("Revenue: 4200000") so a query like "what was Q4 revenue" can match a
    labeled value instead of an anonymous number.
    """

    def parse(self, data: bytes) -> CanonicalDocument:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sections = []
        position = 0

        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            header = [_cell_to_str(c) for c in rows[0]]
            for row in rows[1:]:
                values = [_cell_to_str(c) for c in row]
                if not any(v.strip() for v in values):
                    continue

                pairs = [f"{h}: {v}" for h, v in zip(header, values) if h.strip()]
                text = "; ".join(pairs) if pairs else "; ".join(values)

                sections.append(DocumentSection(
                    text=text,
                    section_type=SectionType.ROW,
                    sheet_name=sheet.title,
                    position=position
                ))
                position += 1

        return CanonicalDocument(sections=sections, source_type="xlsx")


class CsvParser(DocumentParser):
    """Same row/header-pairing approach as XlsxParser, for plain CSV."""

    def parse(self, data: bytes) -> CanonicalDocument:
        text_stream = io.StringIO(data.decode("utf-8", errors="replace"))
        rows = list(csv.reader(text_stream))
        sections = []

        if not rows:
            return CanonicalDocument(sections=[], source_type="csv")

        header = rows[0]
        position = 0
        for row in rows[1:]:
            if not any(cell.strip() for cell in row):
                continue

            pairs = [f"{h}: {v}" for h, v in zip(header, row) if h.strip()]
            text = "; ".join(pairs) if pairs else "; ".join(row)

            sections.append(DocumentSection(text=text, section_type=SectionType.ROW, position=position))
            position += 1

        return CanonicalDocument(sections=sections, source_type="csv")


def _cell_to_str(value) -> str:
    return "" if value is None else str(value)


PARSER_REGISTRY = {
    "pdf": PdfParser(),
    "docx": DocxParser(),
    "xlsx": XlsxParser(),
    "csv": CsvParser(),
}


def get_parser(file_type: str) -> DocumentParser:
    parser = PARSER_REGISTRY.get(file_type)
    if parser is None:
        raise ValueError(f"No parser registered for file type: {file_type!r}")
    return parser
